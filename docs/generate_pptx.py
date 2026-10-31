"""
Spring AI Ops — 프레젠테이션 생성 스크립트 (README 기반 수정판)
색상: Spring Cloud Eureka / aiops.css 기반
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# ── 색상 팔레트 (aiops.css 기반) ──────────────────────────────
C_DARK   = RGBColor(0x34, 0x30, 0x2d)
C_GREEN  = RGBColor(0x6d, 0xb3, 0x3f)
C_GREEN2 = RGBColor(0x5a, 0x96, 0x32)
C_LGREEN = RGBColor(0xf1, 0xf8, 0xe9)
C_WHITE  = RGBColor(0xff, 0xff, 0xff)
C_BG     = RGBColor(0xf5, 0xf5, 0xf5)
C_BORDER = RGBColor(0xdd, 0xdd, 0xdd)
C_TEXT   = RGBColor(0x33, 0x33, 0x33)
C_MUTED  = RGBColor(0x55, 0x55, 0x55)
C_RED    = RGBColor(0xc0, 0x39, 0x2b)
C_BLUE   = RGBColor(0x1a, 0x56, 0xc4)
C_ORANGE = RGBColor(0xe6, 0x75, 0x22)  # Prometheus orange

W = Inches(13.33)
H = Inches(7.5)
DOCS = os.path.dirname(os.path.abspath(__file__))


def rgb(r, g, b):
    return RGBColor(r, g, b)


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=Pt(0)):
    shape = slide.shapes.add_shape(
        1, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = line
        shape.line.width = line_w
    else:
        shape.line.fill.background()
    return shape


def add_text_box(slide, text, x, y, w, h,
                 font_size=Pt(14), bold=False, color=C_TEXT,
                 align=PP_ALIGN.LEFT, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = font_size
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return tb


def add_image(slide, path, x, y, w, h=None):
    if not os.path.exists(path):
        return
    if h:
        slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    else:
        slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w))


def navbar(slide, title="Spring AI Ops", subtitle=""):
    add_rect(slide, 0, 0, 13.33, 0.6, fill=C_DARK)
    add_rect(slide, 0, 0.6, 13.33, 0.05, fill=C_GREEN)
    add_text_box(slide, "🍃 " + title, 0.25, 0.08, 7, 0.5,
                 font_size=Pt(18), bold=True, color=C_WHITE)
    if subtitle:
        add_text_box(slide, subtitle, 6.5, 0.12, 6.6, 0.4,
                     font_size=Pt(11), color=rgb(0xa8, 0xa2, 0x9e),
                     align=PP_ALIGN.RIGHT)


def section_header(slide, text, x=0.3, y=0.85):
    add_rect(slide, x, y, 0.06, 0.3, fill=C_GREEN)
    add_text_box(slide, text, x + 0.12, y - 0.02, 12.5, 0.38,
                 font_size=Pt(16), bold=True, color=C_DARK)


def flow_box(slide, label, x, y, w=2.6, h=0.50,
             bg=C_DARK, text_color=C_WHITE, font_size=Pt(10.5)):
    add_rect(slide, x, y, w, h, fill=bg, line=C_GREEN2, line_w=Pt(1))
    add_text_box(slide, label, x + 0.08, y + 0.04, w - 0.16, h - 0.06,
                 font_size=font_size, bold=True, color=text_color,
                 align=PP_ALIGN.CENTER)


def arrow_down(slide, x, y, color=C_GREEN):
    add_text_box(slide, "↓", x, y, 0.5, 0.3,
                 font_size=Pt(16), bold=True, color=color,
                 align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════
prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
blank = prs.slide_layouts[6]


# ──────────────────────────────────────────────────────────────
# 슬라이드 1: 타이틀
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
add_rect(s, 0, 0, 13.33, 1.7, fill=C_DARK)
add_rect(s, 0, 1.7, 13.33, 0.08, fill=C_GREEN)
add_rect(s, 0, 6.9, 13.33, 0.6, fill=C_DARK)
add_rect(s, 0, 6.9, 13.33, 0.05, fill=C_GREEN)

add_text_box(s, "🍃  Spring AI Ops", 1.2, 0.28, 11, 0.92,
             font_size=Pt(42), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
add_text_box(s, "AI 기반 운영 자동화 플랫폼", 1.2, 1.15, 11, 0.5,
             font_size=Pt(18), color=rgb(0xa8, 0xa2, 0x9e), align=PP_ALIGN.CENTER)

FEATURES = [
    ("🔍", "Code Risk 분석",  "Git 전체 소스 AI 보안·품질 점검\n(Map-Reduce 전략)"),
    ("📝", "Code Review",    "GitHub / GitLab push 자동 리뷰\n(단일 엔드포인트, 자동 분기)"),
    ("🚨", "Error Analysis", "Grafana + Loki + Prometheus\n+ 소스 snippet AI 분석"),
]
card_w = 3.5
gap    = 0.44
sx     = (13.33 - card_w * 3 - gap * 2) / 2

for i, (icon, title, desc) in enumerate(FEATURES):
    cx = sx + i * (card_w + gap)
    add_rect(s, cx, 2.1, card_w, 2.45, fill=C_WHITE, line=C_GREEN, line_w=Pt(1.5))
    add_rect(s, cx, 2.1, card_w, 0.08, fill=C_GREEN)
    add_text_box(s, icon, cx, 2.28, card_w, 0.68,
                 font_size=Pt(34), align=PP_ALIGN.CENTER)
    add_text_box(s, title, cx, 2.93, card_w, 0.42,
                 font_size=Pt(15), bold=True, color=C_DARK, align=PP_ALIGN.CENTER)
    add_text_box(s, desc, cx + 0.1, 3.36, card_w - 0.2, 0.82,
                 font_size=Pt(10.5), color=C_MUTED, align=PP_ALIGN.CENTER)

add_text_box(s,
    "Spring Boot 3.4 · Kotlin · Spring AI 1.1.0 · OpenAI / Anthropic / DeepSeek · Loki · Prometheus",
    1, 4.75, 11.33, 0.38,
    font_size=Pt(11.5), color=C_MUTED, align=PP_ALIGN.CENTER)


# ──────────────────────────────────────────────────────────────
# 슬라이드 2: 프로젝트 개요
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="프로젝트 개요")
section_header(s, "프로젝트 개요")

add_text_box(s,
    "Spring AI Ops는 모니터링 · 형상관리 도구체인을 LLM과 연결하는 AI 기반 운영 자동화 도구입니다.\n"
    "Git 저장소 코드 분석, push 기반 코드 리뷰, Grafana 장애 자동 대응 결과를 WebSocket 대시보드로 실시간 스트리밍합니다.",
    0.4, 1.25, 12.5, 0.65, font_size=Pt(12.5), color=C_MUTED)

rows = [
    ("기능",               "입력 소스",                           "LLM 출력"),
    ("Code Risk 분석",     "Git 저장소 전체 소스\n(단일 / Map-Reduce 자동 선택)",
                                                                "보안 취약점·개선 권고 리포트 + JSON 이슈 목록"),
    ("자동 Code Review",   "GitHub / GitLab push\n(단일 엔드포인트, 자동 분기)",
                                                                "코드 품질·보안·잠재버그·개선 제안 리뷰"),
    ("Error Analysis",    "Grafana Alert + Loki 로그\n+ Prometheus 메트릭(선택) + 소스 snippet",
                                                                "근본원인·영향범위·조치방법 + 소스 수정 권고안"),
]
col_w = [2.8, 4.6, 4.8]
col_x = [0.4, 3.3, 8.0]
row_h = 0.7

for ri, row in enumerate(rows):
    is_hd = ri == 0
    bg = C_DARK if is_hd else (C_LGREEN if ri % 2 == 0 else C_WHITE)
    tc = C_WHITE if is_hd else C_DARK
    for ci, (cell, cw, cx) in enumerate(zip(row, col_w, col_x)):
        add_rect(s, cx, 2.0 + ri * row_h, cw - 0.05, row_h,
                 fill=bg, line=C_BORDER, line_w=Pt(0.5))
        add_text_box(s, cell, cx + 0.1, 2.04 + ri * row_h,
                     cw - 0.25, row_h - 0.06,
                     font_size=Pt(10.5 if not is_hd else 11.5),
                     bold=is_hd, color=tc)

# 주요 기능 카드
add_rect(s, 0.4, 4.9, 12.5, 0.05, fill=C_GREEN)
add_text_box(s, "핵심 설계 포인트", 0.4, 5.02, 5, 0.35,
             font_size=Pt(12.5), bold=True, color=C_DARK)

pts = [
    ("🔄", "동적 LLM 전환",
     "OpenAI / Anthropic / DeepSeek\n재시작 없이 UI에서 즉시 전환"),
    ("🔒", "AES-256-GCM 암호화",
     "API Key · Git Token\nRedis 저장 전 자동 암호화"),
    ("⚡", "Virtual Thread 비동기",
     "Webhook 즉시 반환\nJava 21 가상 스레드로 분석 처리"),
    ("📦", "RDB 미사용",
     "Redis 단일 저장소\nEmbedded Redis 자동 기동"),
]
for i, (icon, title, desc) in enumerate(pts):
    px = 0.4 + i * 3.25
    add_rect(s, px, 5.42, 3.15, 1.82, fill=C_LGREEN,
             line=C_GREEN, line_w=Pt(1))
    add_rect(s, px, 5.42, 3.15, 0.07, fill=C_GREEN)
    add_text_box(s, icon + "  " + title, px + 0.1, 5.56, 2.9, 0.42,
                 font_size=Pt(12), bold=True, color=C_DARK)
    add_text_box(s, desc, px + 0.1, 5.98, 2.9, 1.15,
                 font_size=Pt(10.5), color=C_MUTED)


# ──────────────────────────────────────────────────────────────
# 슬라이드 3: Code Risk — 기능 소개
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="Code Risk 분석")
section_header(s, "🔍  Code Risk 분석 — 기능 소개")

add_text_box(s,
    "대시보드에서 버튼 클릭 한 번으로 등록된 Git 저장소 전체를 AI가 검토하여\n"
    "보안 취약점 · 코드 품질 이슈 · 개선 권고사항을 파일별로 실시간 보고합니다.",
    0.4, 1.3, 12.5, 0.6, font_size=Pt(12.5), color=C_MUTED)

left_items = [
    "JGit으로 저장소 클론 (GitHub / GitLab 토큰 자동 적용)",
    "27,000 토큰 임계값 기준 → 단일 호출 또는 Map-Reduce 자동 선택",
    "Map: 청크 분할 → 최대 3개 병렬 LLM 분석 (호출 간 1,000ms 지연)",
    "Reduce: 청크 결과 통합 → 최종 LLM 호출 (중복 제거 · 전체 요약)",
    "이슈: 파일 · 라인 · 심각도(HIGH/MEDIUM/LOW) · 코드 스니펫 · 권고사항",
    "Persistent Repository Storage 지원 (반복 checkout 비용 절감)",
    "분석 진행 상황 WebSocket 실시간 스트리밍 (/topic/analysis/status)",
    "malformed JSON 복구 파서 — 부분 데이터 최대 복구 (fallback 파서)",
]
add_rect(s, 0.4, 2.0, 6.3, 0.35, fill=C_DARK)
add_text_box(s, "주요 특징", 0.5, 2.02, 5.5, 0.32,
             font_size=Pt(11), bold=True, color=C_WHITE)
for i, item in enumerate(left_items):
    bg = C_LGREEN if i % 2 == 0 else C_WHITE
    add_rect(s, 0.4, 2.37 + i * 0.48, 6.3, 0.46,
             fill=bg, line=C_BORDER, line_w=Pt(0.5))
    add_text_box(s, "▸  " + item, 0.55, 2.4 + i * 0.48, 6.05, 0.42,
                 font_size=Pt(10.5), color=C_TEXT)

# 오른쪽: 심각도 배지
add_rect(s, 7.0, 2.0, 5.9, 0.35, fill=C_DARK)
add_text_box(s, "심각도 분류", 7.1, 2.02, 5.5, 0.32,
             font_size=Pt(11), bold=True, color=C_WHITE)

sev_rows = [
    ("HIGH",   rgb(0xa9, 0x44, 0x42), rgb(0xf2, 0xde, 0xde),
     "보안 취약점, SQL Injection, 민감정보 노출\n→ 즉시 수정 필요"),
    ("MEDIUM", rgb(0x8a, 0x6d, 0x3b), rgb(0xfc, 0xf8, 0xe3),
     "잠재 버그 · 성능 이슈\n→ 우선순위 높게 처리 권고"),
    ("LOW",    rgb(0x66, 0x66, 0x66), rgb(0xf0, 0xf0, 0xf0),
     "코드 스타일 · 가독성 · 마이너 개선\n→ 리팩터링 기회"),
]
for i, (label, tc, bg, desc) in enumerate(sev_rows):
    ry = 2.48 + i * 0.98
    add_rect(s, 7.0, ry, 5.9, 0.88, fill=C_WHITE, line=C_BORDER, line_w=Pt(0.5))
    add_rect(s, 7.0, ry, 0.9, 0.88, fill=bg, line=C_BORDER, line_w=Pt(0.5))
    add_text_box(s, label, 7.0, ry + 0.2, 0.9, 0.5,
                 font_size=Pt(11), bold=True, color=tc, align=PP_ALIGN.CENTER)
    add_text_box(s, desc, 8.0, ry + 0.14, 4.75, 0.65,
                 font_size=Pt(11), color=C_TEXT)

# Persistent Repo 설명
add_rect(s, 7.0, 5.52, 5.9, 0.05, fill=C_GREEN)
add_text_box(s, "Persistent Repository Storage", 7.0, 5.64, 5.9, 0.35,
             font_size=Pt(11), bold=True, color=C_DARK)
add_rect(s, 7.0, 6.02, 5.9, 1.35, fill=C_LGREEN, line=C_GREEN, line_w=Pt(0.5))
add_text_box(s,
    "• repository.stored=true 설정 시 로컬 저장소를 보관\n"
    "• Redis lock 기반 branch sync → 반복 clone 비용 제거\n"
    "• 동기화 실패 시 임시 clone 방식으로 자동 fallback\n"
    "• app 삭제 / 이름 변경 시 저장소 자동 정리",
    7.1, 6.08, 5.7, 1.22, font_size=Pt(10.5), color=C_TEXT)


# ──────────────────────────────────────────────────────────────
# 슬라이드 4: Code Risk — Flow
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="Code Risk 분석")
section_header(s, "🔍  Code Risk 분석 — 처리 Flow")

BW = 2.55
BH = 0.44
CX = 5.4

steps_top = [
    "대시보드 — 분석 실행 버튼 클릭",
    "POST /api/code-risk",
    "JGit으로 저장소 준비\n(persistent sync 또는 임시 clone)",
    "소스 파일 수집 → 토큰 수 추정",
]
split_y_start = 1.1
for i, step in enumerate(steps_top):
    by = split_y_start + i * (BH + 0.14)
    flow_box(s, step, CX - BW/2, by, w=BW, h=BH, bg=C_DARK, font_size=Pt(10))
    if i < len(steps_top) - 1:
        arrow_down(s, CX - 0.25, by + BH + 0.01)

split_y = split_y_start + len(steps_top) * (BH + 0.14) - 0.04
arrow_down(s, CX - 0.25, split_y)
add_text_box(s, "토큰 수에 따라 자동 분기", CX - 2, split_y + 0.31, 4, 0.27,
             font_size=Pt(9.5), color=C_MUTED, align=PP_ALIGN.CENTER, italic=True)

branch_y = split_y + 0.6

# 왼쪽 분기
lx = 1.1
add_rect(s, lx, branch_y, BW, 0.28, fill=C_GREEN)
add_text_box(s, "≤ 27,000 토큰 (단일 호출)", lx, branch_y + 0.03, BW, 0.24,
             font_size=Pt(9.5), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
for i, step in enumerate(["전체 번들 → LLM 단일 호출", "이슈 JSON + Markdown 반환"]):
    by = branch_y + 0.34 + i * (BH + 0.1)
    flow_box(s, step, lx, by, w=BW, h=BH,
             bg=rgb(0x4a, 0x7a, 0x2a) if i == 1 else C_DARK, font_size=Pt(10))
    if i == 0:
        arrow_down(s, lx + BW/2 - 0.25, by + BH)

# 오른쪽 분기
rx = 6.8
rw = 3.5
add_rect(s, rx, branch_y, rw, 0.28, fill=C_RED)
add_text_box(s, "> 27,000 토큰 (Map-Reduce)", rx, branch_y + 0.03, rw, 0.24,
             font_size=Pt(9.5), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)

mr_steps = [
    ("[Map] 청크 분할",                           C_DARK),
    ("병렬 LLM 분석 (최대 3개, 1,000ms 지연)",     rgb(0x8b, 0x3a, 0x28)),
    ("[Reduce] 청크 결과 통합",                    C_DARK),
    ("최종 LLM 호출 → 중복 제거 · 전체 요약",       rgb(0x4a, 0x7a, 0x2a)),
]
for i, (step, bg) in enumerate(mr_steps):
    by = branch_y + 0.34 + i * (BH + 0.1)
    flow_box(s, step, rx, by, w=rw, h=BH, bg=bg, font_size=Pt(10))
    if i < len(mr_steps) - 1:
        arrow_down(s, rx + rw/2 - 0.25, by + BH - 0.04)

# 최종 공통 흐름
final_steps = [
    ("CodeRiskRecord → Redis 저장",              C_DARK),
    ("WebSocket /topic/analysis/result → 브라우저", C_GREEN2),
]
final_y = branch_y + 0.34 + len(mr_steps) * (BH + 0.1) + 0.1
arrow_down(s, CX - 0.25, final_y)
for i, (step, bg) in enumerate(final_steps):
    by = final_y + 0.3 + i * (BH + 0.1)
    flow_box(s, step, CX - BW/2, by, w=BW, h=BH, bg=bg, font_size=Pt(10))
    if i == 0:
        arrow_down(s, CX - 0.25, by + BH)

# 오른쪽 노트
add_rect(s, 10.55, branch_y + 0.32, 2.55, 1.98, fill=C_LGREEN,
         line=C_GREEN, line_w=Pt(0.8))
add_text_box(s, "설계 배경", 10.55, branch_y + 0.34, 2.55, 0.3,
             font_size=Pt(10), bold=True, color=C_DARK, align=PP_ALIGN.CENTER)
add_text_box(s,
    "• 단순 분할 시 파일 간\n  연관성 손실 → 품질 저하\n\n"
    "• Map-Reduce로 청크별\n  분석 후 Reduce 통합\n\n"
    "• Rate Limit 회피:\n  1,000ms 호출 간 지연",
    10.65, branch_y + 0.65, 2.35, 1.6, font_size=Pt(9.5), color=C_TEXT)


# ──────────────────────────────────────────────────────────────
# 슬라이드 5: Code Risk — 스크린샷
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="Code Risk 분석")
section_header(s, "🔍  Code Risk 분석 — 화면")

add_rect(s, 0.3, 1.3, 6.0, 0.3, fill=C_DARK)
add_text_box(s, "분석 진행 상황 (실시간 스트리밍)", 0.4, 1.32, 5.8, 0.27,
             font_size=Pt(10.5), bold=True, color=C_WHITE)
add_image(s, os.path.join(DOCS, "Progress.png"), 0.3, 1.63, 6.0)

add_rect(s, 7.0, 1.3, 6.0, 0.3, fill=C_DARK)
add_text_box(s, "코드 위험 분석 결과 (파일별 이슈 목록)", 7.1, 1.32, 5.8, 0.27,
             font_size=Pt(10.5), bold=True, color=C_WHITE)
add_image(s, os.path.join(DOCS, "CodeRisk.png"), 7.0, 1.63, 6.0)

add_rect(s, 0.3, 5.1, 12.7, 0.3, fill=C_DARK)
add_text_box(s, "소스 코드 수정 제안 (Before / After 좌우 비교 팝업)", 0.4, 5.12, 12.5, 0.27,
             font_size=Pt(10.5), bold=True, color=C_WHITE)
add_image(s, os.path.join(DOCS, "SuggestedSource.png"), 0.3, 5.43, 12.7)


# ──────────────────────────────────────────────────────────────
# 슬라이드 6: Code Review — 기능 소개
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="Code Review")
section_header(s, "📝  Code Review — 기능 소개")

add_text_box(s,
    "GitHub / GitLab push 이벤트를 단일 엔드포인트에서 수신하여 LLM이 변경 파일별 diff를 자동 리뷰합니다.\n"
    "코드 품질 · 잠재 버그 · 보안 취약점 · 개선 제안을 구조화된 보고서로 WebSocket 대시보드에 전달합니다.",
    0.4, 1.3, 12.5, 0.6, font_size=Pt(12.5), color=C_MUTED)

left_items = [
    "GitHub / GitLab 단일 엔드포인트 (/webhook/git/{app})",
    "X-Gitlab-Event 헤더 유무로 플랫폼 자동 분기 처리",
    "GitHub: Compare API (before → after SHA)",
    "GitHub 초기 push (before = 0000...) 케이스 별도 처리",
    "GitLab: Repository Compare / Commit Diff API",
    "변경 파일별 diff → LLM 코드 리뷰",
    "CodeReviewRecord → Redis 저장 (히스토리 관리)",
    "WebSocket /topic/commit → 브라우저 실시간 전달",
]
add_rect(s, 0.4, 2.0, 6.4, 0.35, fill=C_DARK)
add_text_box(s, "주요 특징", 0.5, 2.02, 5.8, 0.32,
             font_size=Pt(11), bold=True, color=C_WHITE)
for i, item in enumerate(left_items):
    bg = C_LGREEN if i % 2 == 0 else C_WHITE
    add_rect(s, 0.4, 2.37 + i * 0.48, 6.4, 0.46,
             fill=bg, line=C_BORDER, line_w=Pt(0.5))
    add_text_box(s, "▸  " + item, 0.55, 2.4 + i * 0.48, 6.15, 0.42,
                 font_size=Pt(10.5), color=C_TEXT)

# 오른쪽: 리뷰 관점
add_rect(s, 7.1, 2.0, 5.85, 0.35, fill=C_DARK)
add_text_box(s, "리뷰 관점", 7.2, 2.02, 5.5, 0.32,
             font_size=Pt(11), bold=True, color=C_WHITE)

review_pts = [
    ("코드 품질",  "가독성, 유지보수성, 명명 규칙, 구조적 문제"),
    ("잠재 버그",  "NPE, 경계 조건, 예외 처리, 동시성 문제"),
    ("보안",      "인증·인가, 입력값 검증, 민감정보 노출"),
    ("개선 제안",  "리팩터링 기회, 성능 최적화, 패턴 적용"),
]
for i, (label, desc) in enumerate(review_pts):
    ry = 2.48 + i * 0.88
    add_rect(s, 7.1, ry, 5.85, 0.78, fill=C_WHITE, line=C_BORDER, line_w=Pt(0.5))
    add_rect(s, 7.1, ry, 1.2, 0.78, fill=C_LGREEN, line=C_GREEN, line_w=Pt(0.5))
    add_text_box(s, label, 7.1, ry + 0.2, 1.2, 0.4,
                 font_size=Pt(11), bold=True, color=C_DARK, align=PP_ALIGN.CENTER)
    add_text_box(s, desc, 8.4, ry + 0.2, 4.45, 0.42,
                 font_size=Pt(11), color=C_TEXT)

# LLM 프롬프트
add_rect(s, 7.1, 6.03, 5.85, 0.05, fill=C_GREEN)
add_text_box(s, "LLM 시스템 역할 프롬프트", 7.1, 6.15, 5.85, 0.32,
             font_size=Pt(11), bold=True, color=C_DARK)
add_rect(s, 7.1, 6.5, 5.85, 0.88, fill=rgb(0x1e, 0x1e, 0x1e))
add_text_box(s,
    '"시니어 소프트웨어 엔지니어로서 코드 diff를 리뷰하고\n'
    ' 코드 품질 · 잠재 버그 · 보안 · 개선 제안을 구조화된 보고서로 제공하십시오."',
    7.2, 6.54, 5.65, 0.8,
    font_size=Pt(10.5), color=rgb(0xce, 0x91, 0x78), italic=True)


# ──────────────────────────────────────────────────────────────
# 슬라이드 7: Code Review — Flow
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="Code Review")
section_header(s, "📝  Code Review — 처리 Flow (GitHub / GitLab)")

BW2 = 3.2
BH2 = 0.46
CX2 = 6.67

top_steps = [
    ("git push 발생",                              C_DARK),
    ("POST /webhook/git/{application}\n수신",       C_DARK),
    ("X-Gitlab-Event 헤더 감지 → 플랫폼 자동 분기", C_DARK),
]
start_y2 = 1.12
for i, (label, bg) in enumerate(top_steps):
    by = start_y2 + i * (BH2 + 0.15)
    flow_box(s, label, CX2 - BW2/2, by, w=BW2, h=BH2, bg=bg, font_size=Pt(10))
    if i < len(top_steps) - 1:
        arrow_down(s, CX2 - 0.25, by + BH2)

split_y2 = start_y2 + len(top_steps) * (BH2 + 0.15) - 0.05
arrow_down(s, CX2 - 0.25, split_y2)

branch_y2 = split_y2 + 0.3
# GitHub 분기 (왼쪽)
lx2 = 1.1
lw2 = 3.2
add_rect(s, lx2, branch_y2, lw2, 0.28, fill=C_BLUE)
add_text_box(s, "GitHub 경로", lx2, branch_y2 + 0.03, lw2, 0.23,
             font_size=Pt(10), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
gh_steps = [
    "owner / repo / SHA 추출",
    "GitHub Compare API 호출\n(0000... 첫 push 케이스 별도 처리)",
    "변경 파일별 diff → LLM 코드 리뷰",
]
for i, step in enumerate(gh_steps):
    by = branch_y2 + 0.34 + i * (BH2 + 0.12)
    flow_box(s, step, lx2, by, w=lw2, h=BH2,
             bg=rgb(0x1a, 0x3d, 0x88) if i < 2 else rgb(0x4a, 0x7a, 0x2a),
             font_size=Pt(10))
    if i < len(gh_steps) - 1:
        arrow_down(s, lx2 + lw2/2 - 0.25, by + BH2)

# GitLab 분기 (오른쪽)
rx2 = 9.0
rw2 = 3.2
add_rect(s, rx2, branch_y2, rw2, 0.28, fill=rgb(0xfc, 0x6d, 0x26))
add_text_box(s, "GitLab 경로", rx2, branch_y2 + 0.03, rw2, 0.23,
             font_size=Pt(10), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
gl_steps = [
    "project / SHA 추출",
    "GitLab Repository API 호출\n(Compare / Commit Diff)",
    "변경 파일별 diff → LLM 코드 리뷰",
]
for i, step in enumerate(gl_steps):
    by = branch_y2 + 0.34 + i * (BH2 + 0.12)
    flow_box(s, step, rx2, by, w=rw2, h=BH2,
             bg=rgb(0xa8, 0x43, 0x00) if i < 2 else rgb(0x4a, 0x7a, 0x2a),
             font_size=Pt(10))
    if i < len(gl_steps) - 1:
        arrow_down(s, rx2 + rw2/2 - 0.25, by + BH2)

# 수렴
merge_y = branch_y2 + 0.34 + len(gh_steps) * (BH2 + 0.12) + 0.1
arrow_down(s, CX2 - 0.25, merge_y)
final2 = [
    ("CodeReviewRecord → Redis 저장", C_DARK),
    ("WebSocket /topic/commit → 브라우저 실시간 전달", C_GREEN2),
]
for i, (step, bg) in enumerate(final2):
    by = merge_y + 0.3 + i * (BH2 + 0.1)
    flow_box(s, step, CX2 - BW2/2, by, w=BW2, h=BH2, bg=bg, font_size=Pt(10))
    if i == 0:
        arrow_down(s, CX2 - 0.25, by + BH2)


# ──────────────────────────────────────────────────────────────
# 슬라이드 8: Code Review — 스크린샷
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="Code Review")
section_header(s, "📝  Code Review — 화면")

add_rect(s, 0.3, 1.3, 5.9, 0.3, fill=C_DARK)
add_text_box(s, "GitHub Push 이벤트 수신", 0.4, 1.32, 5.7, 0.27,
             font_size=Pt(10.5), bold=True, color=C_WHITE)
add_image(s, os.path.join(DOCS, "Push.png"), 0.3, 1.63, 5.9)

add_rect(s, 7.0, 1.3, 6.0, 0.3, fill=C_DARK)
add_text_box(s, "AI 코드 리뷰 결과 (파일별 diff + 리뷰 보고서)", 7.1, 1.32, 5.8, 0.27,
             font_size=Pt(10.5), bold=True, color=C_WHITE)
add_image(s, os.path.join(DOCS, "CodeReview.png"), 7.0, 1.63, 6.0)


# ──────────────────────────────────────────────────────────────
# 슬라이드 9: Error Analysis — 기능 소개
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="Error Analysis")
section_header(s, "🚨  Error Analysis — 기능 소개")

add_text_box(s,
    "Grafana 알림 수신 즉시 Loki 로그 · Prometheus 메트릭(선택) · 소스 저장소를 병렬로 조회하고\n"
    "JVM stack trace에서 소스 snippet을 추출하여 LLM이 근본원인 · 영향범위 · 조치방법을 실시간 보고합니다.",
    0.4, 1.3, 12.5, 0.65, font_size=Pt(12.5), color=C_MUTED)

left_items = [
    "Grafana 알림 수신 즉시 전체 파이프라인 자동 실행",
    "alart 레이블(job · namespace · pod) → Loki 스트림 셀렉터 자동 구성",
    "lokiStartNano(-5분) ~ lokiEndNano 범위 로그 자동 조회",
    "Prometheus 메트릭 병렬 조회 (선택 — prometheus.url 미설정 시 스킵)",
    "소스 저장소 병렬 checkout (deploy branch 설정 시 수행)",
    "Loki 로그에서 JVM stack trace 파싱 → 관련 소스 파일 · 라인 특정",
    "focused source snippet 추출 (전체 저장소 LLM 전달 않음)",
    "소스 수정 권고안 JSON 반환 (filePath · originalCode · suggestionCode)",
]
add_rect(s, 0.4, 2.05, 6.5, 0.35, fill=C_DARK)
add_text_box(s, "주요 특징", 0.5, 2.07, 6.0, 0.32,
             font_size=Pt(11), bold=True, color=C_WHITE)
for i, item in enumerate(left_items):
    bg = C_LGREEN if i % 2 == 0 else C_WHITE
    add_rect(s, 0.4, 2.42 + i * 0.48, 6.5, 0.46,
             fill=bg, line=C_BORDER, line_w=Pt(0.5))
    add_text_box(s, "▸  " + item, 0.55, 2.45 + i * 0.48, 6.25, 0.42,
                 font_size=Pt(10.5), color=C_TEXT)

# 오른쪽: LLM 입력 구성
add_rect(s, 7.2, 2.05, 5.75, 0.35, fill=C_DARK)
add_text_box(s, "LLM 입력 구성 (4가지 컨텍스트)", 7.3, 2.07, 5.5, 0.32,
             font_size=Pt(11), bold=True, color=C_WHITE)

llm_inputs = [
    ("📋", "알림 컨텍스트",    C_RED,    "알림 레이블 (job/namespace/pod)\n심각도 · 발생 시각 · 알림 이름"),
    ("📜", "Loki 로그",       C_DARK,   "ERROR / WARN 로그 라인\n(alert 기준 ±5분 자동 조회)"),
    ("📊", "Prometheus 메트릭",C_ORANGE, "같은 레이블·시간 구간의\n메트릭 시계열 (선택)"),
    ("💾", "소스 snippet",    C_GREEN2, "stack trace → 관련 파일·라인\nfocused snippet만 추출"),
]
for i, (icon, title, color, desc) in enumerate(llm_inputs):
    ry = 2.55 + i * 1.12
    add_rect(s, 7.2, ry, 5.75, 1.02, fill=C_WHITE, line=C_BORDER, line_w=Pt(0.5))
    add_rect(s, 7.2, ry, 5.75, 0.08, fill=color)
    add_text_box(s, f"{icon}  {title}", 7.3, ry + 0.12, 5.55, 0.35,
                 font_size=Pt(11.5), bold=True, color=C_DARK)
    add_text_box(s, desc, 7.3, ry + 0.48, 5.55, 0.48,
                 font_size=Pt(10.5), color=C_MUTED)


# ──────────────────────────────────────────────────────────────
# 슬라이드 10: Error Analysis — Flow
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="Error Analysis")
section_header(s, "🚨  Error Analysis — 처리 Flow (병렬 처리 포함)")

BW3 = 3.0
BH3 = 0.46

# 중앙 컬럼
cx3 = 5.16
top3 = [
    ("Grafana Alert 발생 (FIRING)",              C_RED),
    ("POST /webhook/grafana/{application}",     C_DARK),
    ("RESOLVED 체크 → Early Return",             C_DARK),
    ("알림 레이블 → Loki 셀렉터 · 시간 범위 구성", C_DARK),
]
sy3 = 1.12
for i, (label, bg) in enumerate(top3):
    by = sy3 + i * (BH3 + 0.14)
    flow_box(s, label, cx3 - BW3/2, by, w=BW3, h=BH3, bg=bg, font_size=Pt(10))
    if i < len(top3) - 1:
        arrow_down(s, cx3 - 0.25, by + BH3)

split_y3 = sy3 + len(top3) * (BH3 + 0.14) - 0.04
arrow_down(s, cx3 - 0.25, split_y3)

add_text_box(s, "3가지 병렬 처리", cx3 - 1.5, split_y3 + 0.3, 3.0, 0.28,
             font_size=Pt(10), color=C_MUTED, align=PP_ALIGN.CENTER, italic=True)

par_y = split_y3 + 0.6
PBH = 0.46

# 병렬 1: Loki (왼쪽)
lx3 = 0.25
lw3 = 2.8
add_rect(s, lx3, par_y, lw3, 0.28, fill=C_ORANGE)
add_text_box(s, "① Loki 로그 조회", lx3, par_y + 0.03, lw3, 0.23,
             font_size=Pt(10), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
loki_steps = [
    "GET /loki/api/v1/query_range",
    "ERROR / WARN 로그 수집",
    "JVM stack trace 파싱",
    "소스 파일 resolve\n→ focused snippet 추출",
]
for i, step in enumerate(loki_steps):
    by = par_y + 0.34 + i * (PBH + 0.1)
    flow_box(s, step, lx3, by, w=lw3, h=PBH,
             bg=rgb(0x8b, 0x42, 0x00) if i < 2 else C_DARK,
             font_size=Pt(9.5))
    if i < len(loki_steps) - 1:
        arrow_down(s, lx3 + lw3/2 - 0.25, by + PBH - 0.03)

# 병렬 2: Prometheus (중앙)
px3 = 4.6
pw3 = 2.8
add_rect(s, px3, par_y, pw3, 0.28, fill=C_ORANGE)
add_text_box(s, "② Prometheus 조회 (선택)", px3, par_y + 0.03, pw3, 0.23,
             font_size=Pt(9.5), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
prom_steps = [
    "GET /api/v1/query_range",
    "같은 레이블 · 시간 구간\n메트릭 시계열 수집",
    "prometheus.url 미설정 시\n→ 자동 스킵",
]
for i, step in enumerate(prom_steps):
    by = par_y + 0.34 + i * (PBH + 0.1)
    flow_box(s, step, px3, by, w=pw3, h=PBH,
             bg=rgb(0x8b, 0x42, 0x00) if i == 0 else rgb(0x4a, 0x7a, 0x2a) if i == 2 else C_DARK,
             font_size=Pt(9.5))
    if i < len(prom_steps) - 1:
        arrow_down(s, px3 + pw3/2 - 0.25, by + PBH - 0.03)

# 병렬 3: 소스 저장소 (오른쪽)
sx3 = 9.0
sw3 = 2.8
add_rect(s, sx3, par_y, sw3, 0.28, fill=C_GREEN2)
add_text_box(s, "③ 소스 저장소 준비 (선택)", sx3, par_y + 0.03, sw3, 0.23,
             font_size=Pt(9.5), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
src_steps = [
    "deploy branch 설정 시 수행",
    "JGit checkout\n(persistent sync / 임시 clone)",
    "stack trace 기반\n소스 snippet 추출",
]
for i, step in enumerate(src_steps):
    by = par_y + 0.34 + i * (PBH + 0.1)
    flow_box(s, step, sx3, by, w=sw3, h=PBH,
             bg=C_DARK if i < 2 else rgb(0x4a, 0x7a, 0x2a),
             font_size=Pt(9.5))
    if i < len(src_steps) - 1:
        arrow_down(s, sx3 + sw3/2 - 0.25, by + PBH - 0.03)

# 수렴 → LLM
merge_y3 = par_y + 0.34 + len(loki_steps) * (PBH + 0.1) + 0.08
arrow_down(s, cx3 - 0.25, merge_y3)
final3 = [
    ("알림 + 로그 + 메트릭 + snippet → LLM 분석",    rgb(0x6b, 0x2a, 0x20)),
    ("AnalyzeFiringRecord → Redis 저장",             C_DARK),
    ("WebSocket /topic/firing → 브라우저 실시간 전달", C_GREEN2),
]
for i, (step, bg) in enumerate(final3):
    by = merge_y3 + 0.28 + i * (PBH + 0.1)
    flow_box(s, step, cx3 - BW3/2, by, w=BW3, h=PBH, bg=bg, font_size=Pt(10))
    if i < len(final3) - 1:
        arrow_down(s, cx3 - 0.25, by + PBH - 0.03)


# ──────────────────────────────────────────────────────────────
# 슬라이드 11: Error Analysis — 스크린샷
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="Error Analysis")
section_header(s, "🚨  Error Analysis — 화면")

add_rect(s, 0.3, 1.3, 5.9, 0.3, fill=C_DARK)
add_text_box(s, "Grafana Alerting Webhook 수신 대시보드", 0.4, 1.32, 5.7, 0.27,
             font_size=Pt(10.5), bold=True, color=C_WHITE)
add_image(s, os.path.join(DOCS, "GrafanaAlerting.png"), 0.3, 1.63, 5.9)

add_rect(s, 7.0, 1.3, 6.0, 0.3, fill=C_DARK)
add_text_box(s, "AI 근본원인 분석 결과 (실시간 스트리밍)", 7.1, 1.32, 5.8, 0.27,
             font_size=Pt(10.5), bold=True, color=C_WHITE)
add_image(s, os.path.join(DOCS, "FiringAnalyze.png"), 7.0, 1.63, 6.0)

add_rect(s, 0.3, 5.1, 12.7, 0.3, fill=C_DARK)
add_text_box(s, "소스 수정 권고안 (stack trace 기반 → Before / After 비교 팝업)", 0.4, 5.12, 12.5, 0.27,
             font_size=Pt(10.5), bold=True, color=C_WHITE)
add_image(s, os.path.join(DOCS, "SuggestedSource.png"), 0.3, 5.43, 12.7)


# ──────────────────────────────────────────────────────────────
# 슬라이드 12: Observability & Metrics
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="Observability 연동")
section_header(s, "📡  Observability 연동 — Loki · Prometheus")

add_text_box(s,
    "Loki는 Grafana 장애 분석의 필수 데이터 소스입니다. Prometheus는 선택 사항으로,\n"
    "설정 시 메트릭 시계열을 LLM 분석 컨텍스트에 포함하고 대시보드에 애플리케이션별 시각화를 제공합니다.",
    0.4, 1.3, 12.5, 0.65, font_size=Pt(12.5), color=C_MUTED)

add_rect(s, 0.3, 2.08, 5.9, 0.3, fill=C_DARK)
add_text_box(s, "Observability 연동 설정 화면", 0.4, 2.1, 5.7, 0.27,
             font_size=Pt(10.5), bold=True, color=C_WHITE)
add_image(s, os.path.join(DOCS, "ObservabilityProviders.png"), 0.3, 2.41, 5.9)

add_rect(s, 7.0, 2.08, 6.0, 0.3, fill=C_DARK)
add_text_box(s, "Prometheus 기반 애플리케이션 메트릭 시각화", 7.1, 2.1, 5.8, 0.27,
             font_size=Pt(10.5), bold=True, color=C_WHITE)
add_image(s, os.path.join(DOCS, "MetricVisualization.png"), 7.0, 2.41, 6.0)

# 하단 설명
add_rect(s, 0.3, 6.12, 12.7, 0.05, fill=C_ORANGE)
add_text_box(s, "Prometheus 메트릭 시각화 항목", 0.3, 6.25, 5, 0.32,
             font_size=Pt(11), bold=True, color=C_DARK)
prom_items = [
    "JVM 메모리 사용량 (used% · allocated MB · used MB)",
    "프로세스 업타임 · 열린 파일 수",
    "CPU 사용률 (system · process)",
    "평균 HTTP 레이턴시 · 응답 상태 분포 (2xx/4xx/5xx) — 시계열 차트",
]
for i, item in enumerate(prom_items):
    col = i % 2
    row = i // 2
    add_text_box(s, "▸  " + item, 0.4 + col * 6.5, 6.62 + row * 0.32, 6.3, 0.3,
                 font_size=Pt(10.5), color=C_MUTED)


# ──────────────────────────────────────────────────────────────
# 슬라이드 13: 기술 스택
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
navbar(s, subtitle="기술 스택")
section_header(s, "⚙️  기술 스택")

tech_rows = [
    ("언어",         "Kotlin 2.2 / Java 21",
     "Virtual Thread 기반 비동기 처리"),
    ("프레임워크",    "Spring Boot 3.4.4",
     "WebMVC + WebSocket (STOMP/SockJS) + Embedded Redis"),
    ("AI",           "Spring AI 1.1.0",
     "OpenAI (gpt-4o-mini) / Anthropic (claude-sonnet-4-6) / DeepSeek (deepseek-v4-pro, OpenAI 호환 API)"),
    ("관측성",        "Loki (필수) + Prometheus (선택)",
     "로그 조회 · 메트릭 시계열 · Grafana 알림 연동"),
    ("HTTP 클라이언트","Spring Cloud OpenFeign",
     "Resilience4j Circuit Breaker · TimeLimiter 연동"),
    ("저장소",        "Redis",
     "유일한 데이터 저장소 (RDB 미사용) / 로컬 개발 시 Embedded Redis 자동 기동"),
    ("보안",          "AES-256-GCM",
     "LLM API Key · Git Token 암호화 저장 / CSRF 토큰 보호 (/api/code-risk/**)"),
    ("API 문서",      "springdoc-openapi 2.8.3",
     "Swagger UI (http://localhost:7079/swagger-ui.html)"),
    ("비동기",        "Virtual Thread + Semaphore",
     "SimpleAsyncTaskExecutor (200 동시) + LLM 호출 Semaphore (기본 10)"),
]

add_rect(s, 0.3, 1.3, 12.75, 0.38, fill=C_DARK)
for ci, (label, cw, cx) in enumerate(zip(
        ["구분", "기술", "비고"],
        [2.1, 3.6, 6.85],
        [0.3, 2.5, 6.2])):
    add_text_box(s, label, cx + 0.1, 1.32, cw, 0.35,
                 font_size=Pt(11), bold=True, color=C_WHITE)

for ri, (cat, tech, note) in enumerate(tech_rows):
    bg = C_LGREEN if ri % 2 == 0 else C_WHITE
    ry = 1.7 + ri * 0.49
    add_rect(s, 0.3, ry, 2.15, 0.47, fill=bg, line=C_BORDER, line_w=Pt(0.5))
    add_rect(s, 2.5, ry, 3.65, 0.47, fill=bg, line=C_BORDER, line_w=Pt(0.5))
    add_rect(s, 6.2, ry, 6.85, 0.47, fill=bg, line=C_BORDER, line_w=Pt(0.5))
    add_text_box(s, cat, 0.4, ry + 0.07, 2.0, 0.36,
                 font_size=Pt(10.5), bold=True, color=C_DARK)
    add_text_box(s, tech, 2.6, ry + 0.07, 3.5, 0.36,
                 font_size=Pt(10.5), bold=True, color=C_GREEN2)
    add_text_box(s, note, 6.3, ry + 0.07, 6.65, 0.36,
                 font_size=Pt(10.5), color=C_TEXT)

# 아키텍처 특이사항
add_rect(s, 0.3, 6.15, 12.75, 0.05, fill=C_GREEN)
add_text_box(s, "아키텍처 특이사항", 0.3, 6.27, 5, 0.32,
             font_size=Pt(11.5), bold=True, color=C_DARK)

arch_pts = [
    "Spring AI AutoConfiguration 전면 비활성화 → AiModelService에서 ChatModel 직접 빌드 (재시작 없이 LLM 전환 / DeepSeek는 OpenAI 호환 API 사용)",
    "Virtual Thread + Semaphore: 동시 LLM 호출 수 상한 (기본 10) · cancel-running-future=false로 Virtual Thread interrupt 문제 회피",
]
for i, pt in enumerate(arch_pts):
    add_text_box(s, "▸  " + pt, 0.3, 6.62 + i * 0.32, 12.75, 0.3,
                 font_size=Pt(10.5), color=C_MUTED)


# ──────────────────────────────────────────────────────────────
# 슬라이드 14: Q&A
# ──────────────────────────────────────────────────────────────
s = prs.slides.add_slide(blank)
add_rect(s, 0, 0, 13.33, 7.5, fill=C_WHITE)
add_rect(s, 0, 0, 13.33, 1.8, fill=C_DARK)
add_rect(s, 0, 1.8, 13.33, 0.08, fill=C_GREEN)
add_rect(s, 0, 6.9, 13.33, 0.6, fill=C_DARK)
add_rect(s, 0, 6.9, 13.33, 0.05, fill=C_GREEN)

add_text_box(s, "🍃  Spring AI Ops", 1, 0.28, 11.33, 0.9,
             font_size=Pt(40), bold=True, color=C_WHITE, align=PP_ALIGN.CENTER)
add_text_box(s, "AI 기반 운영 자동화 플랫폼", 1, 1.15, 11.33, 0.5,
             font_size=Pt(18), color=rgb(0xa8, 0xa2, 0x9e), align=PP_ALIGN.CENTER)

summary = [
    ("🔍", "Code Risk 분석",
     "전체 코드베이스 AI 분석\nMap-Reduce · Persistent Repo · 소스 수정 권고"),
    ("📝", "Code Review",
     "GitHub / GitLab push 자동 리뷰\n단일 엔드포인트 · 플랫폼 자동 분기"),
    ("🚨", "Error Analysis",
     "Grafana + Loki + Prometheus\n+ stack trace snippet AI 분석"),
]
card_w3 = 3.6
gap3    = 0.4
sx3b    = (13.33 - card_w3 * 3 - gap3 * 2) / 2

for i, (icon, title, desc) in enumerate(summary):
    cx3b = sx3b + i * (card_w3 + gap3)
    add_rect(s, cx3b, 2.25, card_w3, 2.55, fill=C_WHITE,
             line=C_GREEN, line_w=Pt(1.5))
    add_rect(s, cx3b, 2.25, card_w3, 0.08, fill=C_GREEN)
    add_text_box(s, icon, cx3b, 2.45, card_w3, 0.68,
                 font_size=Pt(34), align=PP_ALIGN.CENTER)
    add_text_box(s, title, cx3b, 3.1, card_w3, 0.42,
                 font_size=Pt(15), bold=True, color=C_DARK, align=PP_ALIGN.CENTER)
    add_text_box(s, desc, cx3b + 0.1, 3.53, card_w3 - 0.2, 0.95,
                 font_size=Pt(10.5), color=C_MUTED, align=PP_ALIGN.CENTER)

add_text_box(s, "Q & A", 0, 5.0, 13.33, 0.9,
             font_size=Pt(38), bold=True, color=C_DARK, align=PP_ALIGN.CENTER)
add_text_box(s, "감사합니다", 0, 5.88, 13.33, 0.55,
             font_size=Pt(20), color=C_MUTED, align=PP_ALIGN.CENTER)


# ── 저장 ──────────────────────────────────────────────────────
output_path = os.path.join(DOCS, "spring-ai-ops-presentation.pptx")
prs.save(output_path)
print(f"✅  저장 완료: {output_path}")
print(f"    슬라이드 수: {len(prs.slides)}장")
